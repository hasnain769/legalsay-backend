from google.adk.agents.llm_agent import LlmAgent
from google.adk.planners import BuiltInPlanner
from google.genai.types import ThinkingConfig
from pydantic import BaseModel, Field
import docx
from docx.shared import RGBColor
from io import BytesIO

class RedlineOutput(BaseModel):
    compliant_text: str = Field(description="The rewritten, legally compliant version of the clause.")
    explanation: str = Field(description="Brief explanation of why this change was made.")

class RedlineAgent(LlmAgent):
    def __init__(self):
        super().__init__(
            model='gemini-2.0-flash',
            name='redline_agent',
            instruction="""You are a precise legal drafter.
            Your task is to rewrite a specific contract clause to be compliant with the given Jurisdiction and favorable to the Recipient.
            
            You will receive:
            1. The Original Clause (Bad Text).
            2. The Jurisdiction.
            3. The Risk/Reason for change.
            
            Your Output:
            - compliant_text: The exact new wording. Keep it minimal and standard.
            - explanation: A 1-sentence explanation (e.g., "Changed Net 60 to Net 30 to improve cash flow.").
            """,
            planner=BuiltInPlanner(thinking_config=ThinkingConfig(include_thoughts=True)),
            output_schema=RedlineOutput,
        )

class DocxRedliner:
    @staticmethod
    def redline_docx(file_content: bytes, original_text: str, new_text: str) -> BytesIO:
        """
        Locates original_text in the docx and replaces it with a redlined version.
        Returns a BytesIO of the modified docx.
        """
        try:
            doc = docx.Document(BytesIO(file_content))
            found = False
            
            for para in doc.paragraphs:
                if original_text in para.text:
                    # Found the paragraph containing the text
                    # We will rebuild the paragraph to include the redline
                    
                    # Split the text to isolate the part to replace
                    # Note: This is a simple split. If the text appears multiple times in the same paragraph,
                    # this logic might need refinement, but it works for MVP.
                    parts = para.text.split(original_text, 1) # Split only on first occurrence
                    
                    # Clear existing runs
                    para.clear()
                    
                    # Rebuild:
                    # 1. Text before
                    if parts[0]:
                        para.add_run(parts[0])
                    
                    # 2. The "Deleted" text (Strikethrough, Red)
                    run_old = para.add_run(original_text)
                    run_old.font.strike = True
                    run_old.font.color.rgb = RGBColor(255, 0, 0) # Red
                    
                    # 3. The "Inserted" text (Underline, Blue)
                    # Add a space before/after if needed, or rely on new_text having it
                    run_new = para.add_run(f" {new_text} ")
                    run_new.font.underline = True
                    run_new.font.color.rgb = RGBColor(0, 0, 255) # Blue
                    
                    # 4. Text after
                    if len(parts) > 1 and parts[1]:
                        para.add_run(parts[1])
                    
                    found = True
                    break # Stop after first match
            
            if not found:
                print(f"WARNING: Could not find exact text match for redlining: '{original_text[:20]}...'")
                # Fallback: Maybe append to end of doc? Or just do nothing.
                # For now, let's just return the original if not found.
            
            output = BytesIO()
            doc.save(output)
            output.seek(0)
            return output
            
        except Exception as e:
            print(f"ERROR: DocxRedliner failed: {e}")
            # Return original content on error to be safe
            return BytesIO(file_content)
